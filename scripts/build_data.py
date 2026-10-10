#!/usr/bin/env python3
"""Gera site/data/summary.json (somente valores agregados) a partir da tabela
TRATATIVAS exportada do Power BI / Forms.

Uso:
    python scripts/build_data.py TRATATIVAS.xlsx --clientes "Consulta Cliente.xlsx"

Aceita .xlsx ou .csv. Nada que identifique cliente ou matrícula é gravado.

Regras, portadas das medidas do modelo Power BI (as quatro tabelas do relatório):

  Economia  (colunas DE: e PARA:, ex.: "2 RESIDENCIAL" -> "3 RESIDENCIAL")
      diferença = economias PARA - economias DE
      Incremento de economia  (diferença > 0)  -> medida Valor_Incremento
      Decremento de economia  (diferença < 0)  -> medida Valor Decremento
      mesma categoria : diferença x tarifa
      categoria muda  : incremento -> diferença x (tarifa PARA - tarifa DE)
                        decremento -> diferença x (tarifa DE - tarifa PARA)
  Categoria (colunas ANTERIOR, ATUAL, QUANTIDADE) -> medida Valor_Incremento_Cat
      QUANTIDADE x (tarifa ATUAL - tarifa ANTERIOR); positivo = incremento, negativo = decremento.
      Categoria sem tarifa ("Outros") vale R$ 0, como no modelo.

  Resultado do mês = soma dos valores das quatro tabelas.

  Além do relatório (cálculo "corrigido", o padrão; --como-powerbi desliga tudo isto):
    * Tarifas oficiais (Comercial 443,56, Industrial 613,16, Pública 129,15) e comércio popular a 60,24.
    * CADÚNICO é Social e "sem fins lucrativos" é Pública.
    * QUANTIDADE escrita como texto ("DE 1 RES. P/ 1 RES. E 2 COM.") é lida como antes -> depois e entra
      como troca de categoria + economia nova/retirada. O mesmo vale para campos DE:/PARA: com duas categorias.
    * Linhas com o mesmo Id e todas as colunas iguais (duplicadas na exportação) contam uma vez.
  Água + esgoto: o valor é multiplicado por 2 nas ligações cujo TIPO_FATURAMENTO da base de clientes é
  "AGUA E ESGOTO" (cruzamento por matrícula). Se a base não tiver essa coluna, vale a regra por município
  (DOUBLE_CITIES). O Power BI atual não faz esse cruzamento: sem --clientes o resultado é o dele.
"""

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict, namedtuple
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "site" / "data" / "summary.json"

# Tarifa mensal por economia, em centavos (valores oficiais).
TARIFFS_CENTS = {
    "Residencial": 8541,
    "Comercial": 44356,
    "Industrial": 61316,
    "Pública": 12915,
    "Pequeno comércio": 22178,
    "Social": 3012,
    "Comércio popular": 6024,
}
# A tabela "Tarifas" do modelo Power BI tem 1 centavo a mais em Comercial, Industrial e Pública.
# Só o modo --como-powerbi a usa, para reproduzir o relatório ao centavo.
POWERBI_TARIFFS_CENTS = {**TARIFFS_CENTS, "Comercial": 44357, "Industrial": 61317, "Pública": 12916}
OTHER = "Outros"  # categoria não reconhecida: sem tarifa, vale R$ 0 (igual ao modelo)

# Localidades que hoje cobram água + esgoto (tarifa dobra). Sem acento, maiúsculas.
DOUBLE_CITIES = ("CORDEIRO", "MIRACEMA", "APERIBE")
UNKNOWN_CITY = "Não identificada"

# double: True/False = tipo de faturamento da ligação; None = a base não informa (vale a regra por município)
Client = namedtuple("Client", "city double")


def tariff(category, powerbi=False):
    return (POWERBI_TARIFFS_CENTS if powerbi else TARIFFS_CENTS).get(category, 0)


# ---------------------------------------------------------------- texto
def plain(value):
    text = unicodedata.normalize("NFKD", "" if value is None else str(value))
    return "".join(c for c in text if not unicodedata.combining(c)).upper().replace("\xa0", " ")


def header_key(name):
    return re.sub(r"[^A-Z0-9]", "", plain(name))


def clean(value):
    return re.sub(r"\s+", " ", plain(value).replace(".", " ").replace(";", " ")).strip()


def category_of(text, powerbi=False, path="categoria"):
    """Texto livre -> categoria de tarifa, ou 'Outros' quando não reconhecido.

    powerbi=True reproduz o relatório: 'popular' vira Comercial (o texto contém "COM") e, nas trocas
    de categoria (ANTERIOR/ATUAL), 'Pública' não acha tarifa e vale R$ 0."""
    s = clean(text)
    if not powerbi:
        if re.search(r"CAD ?UNICO", s):  # tarifa social do CadÚnico
            return "Social"
        if "LUCRATIV" in s:  # entidade sem fins lucrativos paga a tarifa pública
            return "Pública"
    if "POPULAR" in s:
        return "Comercial" if powerbi else "Comércio popular"
    if "PEQUENO" in s or re.search(r"\bP ?COM", s):
        return "Pequeno comércio"
    if re.search(r"\bSOC", s):
        return "Social"
    if re.search(r"\bIND", s):
        return "Industrial"
    if re.search(r"\bPUB", s):
        return OTHER if powerbi and path == "categoria" else "Pública"
    if re.search(r"\bCOM", s):
        return "Comercial"
    if re.search(r"\bRES", s):
        return "Residencial"
    return OTHER


FILLER = re.compile(r"\b(ECONOMIAS?|UNIDADES?|UND|DE|E)\b")


def parse_units(text, fallback=None, powerbi=False):
    """'1 RES E 2 COM' -> Counter(Residencial=1, Comercial=2). None se não houver número.

    Contagem sem categoria entra com a chave None (a categoria sai do outro lado do DE:/PARA:)."""
    s = clean(text)
    numbers = list(re.finditer(r"\d+", s))
    if not numbers:
        return None
    units = Counter()
    for i, m in enumerate(numbers):
        end = numbers[i + 1].start() if i + 1 < len(numbers) else len(s)
        label = FILLER.sub(" ", s[m.end():end]).strip()
        if label:
            units[category_of(label, powerbi, "economia")] += int(m.group())
        elif fallback:
            units[fallback] += int(m.group())
        else:
            units[None] += int(m.group())
    return Counter({k: v for k, v in units.items() if v})


def parse_side(text, fallback=None, powerbi=False):
    """'2 RESIDENCIAL' -> (2, 'Residencial'). '0' -> (0, None).

    Retorna None quando não há número ou quando o campo mistura categorias ('1 RES E 1 COM'),
    porque as medidas do modelo só valorizam uma categoria por lado."""
    units = parse_units(text, fallback, powerbi)
    if units is None or len(units) > 1:
        return None
    if not units:
        return 0, None
    (category, n), = units.items()
    return n, category


def parse_narrative(text):
    """'DE 1 RES. SOCIAL P/ 2 RES. NORMAIS' -> (antes, depois) como Counters por categoria, ou None."""
    s = re.sub(r"^(ALT\s+)?DE\s+", "", clean(text))
    parts = re.split(r"\s*\bP/+\s*|\s+PARA\s+", s, maxsplit=1)
    if len(parts) != 2:
        return None
    before, after = (parse_units(re.sub(r"^DE\s+", "", part)) for part in parts)
    for units in (before, after):  # sem número, sem categoria ou categoria desconhecida: não dá para valorar
        if not units or None in units or OTHER in units:
            return None
    return before, after


def decompose(before, after):
    """Antes/depois por categoria -> movimentos: troca (A->B), economia nova (+B) ou retirada (-A).

    Unidades iguais nos dois lados se anulam; as que sobram em ambos viram trocas de categoria; o excedente
    é economia nova ou retirada. A soma dos valores é sempre (depois x tarifa) - (antes x tarifa)."""
    old, new = Counter(before), Counter(after)
    for c in TARIFFS_CENTS:
        shared = min(old[c], new[c])
        old[c] -= shared
        new[c] -= shared
    by_tariff = sorted(TARIFFS_CENTS, key=TARIFFS_CENTS.get)
    old_left = [c for c in by_tariff for _ in range(old[c])]
    new_left = [c for c in by_tariff for _ in range(new[c])]
    moves = [("troca", a, b, 1) for a, b in zip(old_left, new_left)]
    moves += [("retirada", a, None, 1) for a in old_left[len(new_left):]]
    moves += [("nova", None, b, 1) for b in new_left[len(old_left):]]
    return moves


def move_effects(moves):
    """Movimentos -> efeitos no formato das tabelas: ('categoria'|'economia', de, para, quantidade, valor)."""
    effects = []
    for kind, a, b, qty in moves:
        if kind == "troca":
            effects.append(("categoria", a, b, qty, qty * (tariff(b) - tariff(a))))
        elif kind == "nova":
            effects.append(("economia", b, b, qty, qty * tariff(b)))
        else:
            effects.append(("economia", a, a, -qty, -qty * tariff(a)))
    return effects


def to_number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if "," in text:  # formato brasileiro: 1.234,5
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_date(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return datetime(1899, 12, 30) + timedelta(days=float(value))
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


# ---------------------------------------------------------------- efeitos
def economy_effect(rec, powerbi=False):
    """DE:/PARA: -> ('economia', de, para, diferença, valor em centavos) | None | ('erro', motivo)."""
    de, para = str(rec.get("DE") or "").strip(), str(rec.get("PARA") or "").strip()
    if not (de or para):
        return None
    if not (de and para):
        return "erro", "incompleto"
    fallback = (category_of(rec.get("TIPODEECONOMIA"), powerbi, "economia")
                if str(rec.get("TIPODEECONOMIA") or "").strip() else None)
    if not powerbi:
        story = parse_narrative(de) or parse_narrative(para)  # "DE 5 RES. P/ 4 RES." digitado nos campos
        if story:
            return "varios", move_effects(decompose(*story)), "narrativa"
        mixed_before, mixed_after = parse_units(de, fallback), parse_units(para, fallback)
        if mixed_before is not None and mixed_after is not None and max(len(mixed_before), len(mixed_after)) > 1:
            if None in mixed_before or None in mixed_after or OTHER in mixed_before or OTHER in mixed_after:
                return "erro", "nao_reconhecido"
            return "varios", move_effects(decompose(mixed_before, mixed_after)), "misto"
    before, after = parse_side(de, fallback, powerbi), parse_side(para, fallback, powerbi)
    if before is None or after is None:
        return "erro", "nao_reconhecido"
    (n, a), (m, b) = before, after
    a, b = a or b or OTHER, b or a or OTHER
    diff = m - n
    if diff == 0:
        return "mesma_quantidade", a, b  # o modelo valoriza em R$ 0 e não lista
    if powerbi and diff > 0 and plain(rec.get("QUALFOIAALTERACAODEECONOMIA")).strip() != "INCREMENTO":
        return "marcacao_divergente", a, b  # o relatório só lista incremento marcado como "Incremento" no Forms
    if a == b:
        value = diff * tariff(b, powerbi)
    elif diff > 0:
        value = diff * (tariff(b, powerbi) - tariff(a, powerbi))
    else:
        value = diff * (tariff(a, powerbi) - tariff(b, powerbi))
    return "economia", a, b, diff, value


def category_effect(rec, powerbi=False):
    """ANTERIOR/ATUAL x QUANTIDADE -> ('categoria', anterior, atual, quantidade, valor) | None | ('erro', motivo)."""
    ant, atu = str(rec.get("ANTERIOR") or "").strip(), str(rec.get("ATUAL") or "").strip()
    if not (ant or atu):
        return None
    if not (ant and atu):
        return "erro", "incompleto"
    qty = to_number(rec.get("QUANTIDADE"))
    if qty is None or qty <= 0 or qty != int(qty):
        story = None if powerbi else parse_narrative(rec.get("QUANTIDADE"))  # quantidade escrita como texto
        if story:
            return "varios", move_effects(decompose(*story)), "narrativa"
        return "erro", "quantidade_invalida"
    a, b = category_of(ant, powerbi), category_of(atu, powerbi)
    if a == b:
        return None
    return "categoria", a, b, int(qty), int(qty) * (tariff(b, powerbi) - tariff(a, powerbi))


def row_effects(rec, powerbi=False):
    """Linha -> (efeitos, motivo de pendência, nota).

    nota: 'troca_sem_variacao' (DE:/PARA: com categoria diferente e mesma quantidade, valor R$ 0),
    'marcacao_divergente' (modo --como-powerbi: aumento de economias marcado como "Decremento" no Forms),
    'narrativa' (quantidade ou DE:/PARA: escritos como "DE ... P/ ...") ou 'misto' (duas categorias no campo)."""
    effects, note = [], None
    for found in (economy_effect(rec, powerbi), category_effect(rec, powerbi)):
        if found is None:
            continue
        if found[0] == "erro":
            return [], found[1], None
        if found[0] == "mesma_quantidade":
            note = "troca_sem_variacao" if found[1] != found[2] else None
            continue
        if found[0] == "marcacao_divergente":
            note = "marcacao_divergente"
            continue
        if found[0] == "varios":
            effects.extend(found[1])
            note = found[2]
            continue
        effects.append(found)
    return effects, None, note


# ---------------------------------------------------------------- leitura
def read_table(path, sheet=None):
    """Retorna (colunas_originais, [dict com chaves normalizadas])."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Instale o leitor de Excel: python -m pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        first = text.splitlines()[0] if text else ""
        delimiter = max(",;\t", key=first.count)
        grid = list(csv.reader(text.splitlines(), delimiter=delimiter))

    for start, row in enumerate(grid[:10]):
        if any(header_key(c) in ("HORADECONCLUSAO", "NDALIGACAO", "LIGACAO") for c in row if c):
            break
    else:
        start = 0
    names = [str(c).strip() if c is not None else "" for c in grid[start]]
    keys = [header_key(n) for n in names]
    rows = []
    for line in grid[start + 1:]:
        if not any(c not in (None, "") for c in line):
            continue
        rows.append({k: v for k, v in zip(keys, line) if k})
    return names, rows


def digits(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value or "")).lstrip("0")


def month_key(value):
    """'09/2026', '2026-09-01' ou data -> '2026-09' (vazio se não der para ler)."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    m = re.fullmatch(r"(\d{1,2})/(\d{4})", text)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.match(r"(\d{4})-(\d{2})", text)
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def load_clients(path, col_ligacao=None, col_localidade=None, col_faturamento=None):
    """Base de clientes -> {matrícula: Client(cidade, faz água e esgoto?)}, pela linha do mês mais recente."""
    names, rows = read_table(path)
    keys = {header_key(n): n for n in names if n}

    def pick(wanted, candidates, label, required=True):
        if wanted:
            if header_key(wanted) not in keys:
                sys.exit(f"Coluna '{wanted}' não existe na base de clientes. Colunas: {', '.join(names)}")
            return header_key(wanted)
        for c in candidates:
            for k in keys:
                if c in k and not k.startswith("TIPO"):
                    return k
        if required:
            sys.exit(f"Não achei a coluna de {label} na base de clientes. Use a opção correspondente. "
                     f"Colunas: {', '.join(names)}")
        return None

    k_id = pick(col_ligacao, ("NUMLIGACAO", "LIGACAO", "MATRICULA"), "ligação/matrícula")
    k_city = pick(col_localidade, ("CIDADE", "LOCALIDADE", "MUNICIPIO"), "localidade")
    k_bill = (pick(col_faturamento, (), "tipo de faturamento") if col_faturamento
              else next((k for k in keys if "FATURAMENTO" in k), None))
    k_month = next((k for k in keys if "MESANO" in k), None)

    latest = {}
    for rec in rows:
        ident, city = digits(rec.get(k_id)), str(rec.get(k_city) or "").strip()
        if not ident or not city:  # rodapés e linhas em branco
            continue
        month = month_key(rec.get(k_month)) if k_month else ""
        if ident not in latest or month >= latest[ident][0]:
            double = ("ESGOTO" in plain(rec.get(k_bill))) if k_bill else None
            latest[ident] = (month, Client(city, double))
    return {ident: client for ident, (_, client) in latest.items()}


# ---------------------------------------------------------------- agregação
def build(rows, places=None, only_front=None, since=None, until=None, powerbi=False):
    agg = defaultdict(Counter)
    pending = Counter()
    notes = Counter()
    coverage = Counter()
    dates, top = [], []
    first_by_id = {}
    for number, rec in enumerate(rows, start=2):
        rid = rec.get("ID")
        if not powerbi and rid not in (None, ""):  # linha idêntica repetida na exportação conta uma vez
            if first_by_id.setdefault(rid, rec) is not rec and first_by_id[rid] == rec:
                notes["id_repetido"] += 1
                continue
        if only_front and plain(rec.get("FRENTEDESERVICO")).strip() != plain(only_front).strip():
            continue
        when = parse_date(rec.get("HORADECONCLUSAO"))
        effects, reason, note = row_effects(rec, powerbi)
        if when is None:
            if effects or reason:
                pending[("", UNKNOWN_CITY, "sem_data")] += 1
            continue
        month = when.strftime("%Y-%m")
        if (since and month < since) or (until and month > until):
            continue
        dates.append(when)
        client = places.get(digits(rec.get("MATRICULASDIGITO"))) if places is not None else None
        if isinstance(client, str):  # só a cidade (testes e bases antigas)
            client = Client(client, None)
        city = client.city if client else UNKNOWN_CITY
        if reason:
            pending[(month, city, reason)] += 1
            continue
        if note:
            notes[note] += 1
        if client is None:
            factor = 1  # sem localidade: só água
        elif client.double is not None:
            factor = 2 if client.double else 1
        else:
            factor = 2 if plain(client.city).strip() in DOUBLE_CITIES else 1
        if effects and places is not None:
            coverage["total"] += 1
            coverage["found"] += client is not None
            coverage["double"] += factor == 2
        row_value = 0
        for kind, old, new, qty, value in effects:
            direction = "inc" if (qty > 0 if kind == "economia" else value > 0) else "dec"
            a = agg[(month, city, factor, kind, direction, old, new)]
            a["qty"] += qty
            a["cents"] += value
            a["rows"] += 1
            row_value += value
            notes["outros"] += OTHER in (old, new)
        if effects:
            top.append((abs(row_value * factor), number, month, city, row_value * factor))

    lines = [dict(month=m, city=c, factor=f, kind=k, dir=d, **{"from": o}, to=n,
                  qty=int(a["qty"]), cents=int(a["cents"]), rows=int(a["rows"]))
             for (m, c, f, k, d, o, n), a in sorted(agg.items())]
    out_pending = [dict(month=m, city=c, reason=r, count=n) for (m, c, r), n in sorted(pending.items())]
    feed = {
        "sample": False,
        "source": {
            "tratativas": len(rows),
            "firstDate": min(dates).date().isoformat() if dates else None,
            "lastDate": max(dates).date().isoformat() if dates else None,
            "clientsLinked": places is not None,
            "billing": ("ligacao" if any(getattr(c, "double", None) is not None for c in places.values())
                        else "municipio") if places else None,
            "coverage": dict(found=coverage["found"], total=coverage["total"]) if places is not None else None,
            "rules": "powerbi" if powerbi else "corrigida",
        },
        "tariffsCents": POWERBI_TARIFFS_CENTS if powerbi else TARIFFS_CENTS,
        "doubleCities": sorted(DOUBLE_CITIES),
        "months": sorted({ln["month"] for ln in lines}),
        "lines": lines,
        "pending": out_pending,
    }
    notes.update({f"cobertura_{k}": v for k, v in coverage.items()})
    return feed, sorted(top, reverse=True), notes


def report(feed, top, notes):
    print(f"\n{'mês':8} {'novas':>6} {'retir.':>6} {'trocas↑':>8} {'trocas↓':>8} {'incremento':>12} {'decremento':>12} "
          f"{'líquido s/2×':>13} {'líquido':>12}")
    for month in feed["months"]:
        t = Counter()
        for ln in feed["lines"]:
            if ln["month"] != month:
                continue
            up = ln["dir"] == "inc"
            if ln["kind"] == "economia":
                t["novas" if up else "retir"] += abs(ln["qty"])
            else:
                t["up" if up else "down"] += ln["qty"]
            t["inc" if up else "dec"] += ln["cents"] * ln["factor"]
            t["base"] += ln["cents"]
        print(f"{month:8} {t['novas']:6} {t['retir']:6} {t['up']:8} {t['down']:8} {t['inc'] / 100:12,.2f} "
              f"{t['dec'] / 100:12,.2f} {t['base'] / 100:13,.2f} {(t['inc'] + t['dec']) / 100:12,.2f}")
    print("\n'líquido s/2×' é o resultado sem a cobrança de água + esgoto: deve bater com o Power BI atual.")
    pend = Counter()
    for p in feed["pending"]:
        pend[p["reason"]] += p["count"]
    if pend:
        print("\nPendências (fora dos valores):", dict(pend))
    if notes["cobertura_total"]:
        found, total = notes["cobertura_found"], notes["cobertura_total"]
        print(f"Base de clientes: {found} de {total} tratativas com valor encontradas ({100 * found / total:.1f}%); "
              f"{notes['cobertura_double']} com água e esgoto (2×). As não encontradas foram calculadas só com água.")
    for key, text in (("narrativa", "Quantidade ou DE:/PARA: escritos como texto, lidos como antes → depois"),
                      ("misto", "Campos DE:/PARA: com duas categorias, valorados pela diferença por categoria"),
                      ("id_repetido", "Linhas idênticas (mesmo Id) contadas uma vez")):
        if notes[key]:
            print(f"{text}: {notes[key]}")
    if notes["outros"]:
        print(f"Categorias não reconhecidas, tratadas como 'Outros' (R$ 0): {notes['outros']} movimentos")
    if notes["troca_sem_variacao"]:
        print(f"Linhas DE:/PARA: com categoria diferente e mesma quantidade (o modelo valoriza em R$ 0): "
              f"{notes['troca_sem_variacao']}")
    if notes["marcacao_divergente"]:
        print(f"Aumentos de economia marcados como 'Decremento' no Forms, fora do relatório: {notes['marcacao_divergente']}")
    if top:
        print("\nMaiores impactos individuais (linha da planilha, valor/mês) — confira se não são erro de digitação:")
        for _, number, month, city, signed in top[:5]:
            print(f"  linha {number}: {month} {city}: R$ {signed / 100:,.2f}")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("tratativas", type=Path, help="Exportação da tabela TRATATIVAS (.xlsx ou .csv)")
    p.add_argument("--aba", help="Nome da aba do Excel (padrão: primeira)")
    p.add_argument("--clientes", type=Path, help="Base de clientes para obter a localidade (cruzamento por matrícula)")
    p.add_argument("--col-ligacao", help="Coluna da ligação/matrícula na base de clientes")
    p.add_argument("--col-localidade", help="Coluna da localidade na base de clientes")
    p.add_argument("--col-faturamento", help="Coluna do tipo de faturamento (AGUA / AGUA E ESGOTO) na base de clientes")
    p.add_argument("--frente", help="Considerar só esta FRENTE DE SERVIÇO (ex.: CADASTRO)")
    p.add_argument("--de", dest="since", help="Primeiro mês, AAAA-MM")
    p.add_argument("--ate", dest="until", help="Último mês, AAAA-MM")
    p.add_argument("--como-powerbi", action="store_true",
                   help="Reproduz as particularidades do relatório (comércio popular = comercial, "
                        "pública sem tarifa nas trocas, incremento só se marcado no Forms) para conferir os valores")
    p.add_argument("--exemplo", action="store_true", help="Marca o resultado como dados de exemplo")
    p.add_argument("--saida", type=Path, default=OUTPUT)
    args = p.parse_args()

    names, rows = read_table(args.tratativas, args.aba)
    keys = {header_key(n) for n in names}
    if "HORADECONCLUSAO" not in keys:
        sys.exit(f"Coluna 'Hora de conclusão' não encontrada. Colunas: {', '.join(names)}")
    if not ({"DE", "PARA"} <= keys or {"ANTERIOR", "ATUAL"} <= keys):
        sys.exit("A tabela precisa ter DE: e PARA: (economias) e/ou ANTERIOR e ATUAL (categoria).")
    if args.clientes and "MATRICULASDIGITO" not in keys:
        sys.exit(f"Coluna 'MATRICULA S/ DIGITO' não encontrada. Colunas: {', '.join(names)}")

    places = (load_clients(args.clientes, args.col_ligacao, args.col_localidade, args.col_faturamento)
              if args.clientes else None)
    feed, top, notes = build(rows, places, args.frente, args.since, args.until, args.como_powerbi)
    feed["sample"] = args.exemplo
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    args.saida.write_text(json.dumps(feed, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    report(feed, top, notes)
    print(f"\n{args.saida}: {len(feed['lines'])} linhas agregadas de {len(rows)} tratativas.")


if __name__ == "__main__":
    main()
